from __future__ import annotations
import csv
import io
import json
import re
from pathlib import Path
from typing import Any
from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader

ALLOWED_EXTENSIONS = {".txt", ".md", ".log", ".json", ".jsonl", ".csv", ".html", ".htm", ".pdf", ".docx"}


def _records_to_text(value: Any) -> str:
    """Convert JSON-like incident records to labelled, retrieval-friendly text."""
    if isinstance(value, list):
        records = value
    elif isinstance(value, dict):
        # Common crawler exports wrap record arrays under one of these keys.
        records = None
        for key in ("incidents", "results", "items", "records", "data", "vulnerabilities", "cves"):
            if isinstance(value.get(key), list):
                records = value[key]
                break
        if records is None:
            records = [value]
    else:
        records = [value]

    lines: list[str] = []
    for idx, record in enumerate(records, start=1):
        lines.append(f"Record {idx}:")
        if isinstance(record, dict):
            for key, val in record.items():
                if isinstance(val, (dict, list)):
                    val = json.dumps(val, ensure_ascii=False, sort_keys=True)
                if val is not None and str(val).strip():
                    lines.append(f"{key}: {val}")
        else:
            lines.append(str(record))
        lines.append("")
    return "\n".join(lines).strip()


def extract_text(filename: str, raw: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        allowed = ", ".join(sorted(ALLOWED_EXTENSIONS))
        raise ValueError(f"Unsupported file type '{suffix or '(no extension)'}'. Allowed types: {allowed}")

    if suffix in {".txt", ".md", ".log"}:
        text = raw.decode("utf-8-sig", errors="replace")
    elif suffix == ".json":
        data = json.loads(raw.decode("utf-8-sig", errors="replace"))
        text = _records_to_text(data)
    elif suffix == ".jsonl":
        records = []
        for line_no, line in enumerate(raw.decode("utf-8-sig", errors="replace").splitlines(), start=1):
            if line.strip():
                try:
                    records.append(json.loads(line))
                except json.JSONDecodeError as exc:
                    raise ValueError(f"Invalid JSONL on line {line_no}: {exc.msg}") from exc
        text = _records_to_text(records)
    elif suffix == ".csv":
        decoded = raw.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(decoded))
        records = list(reader)
        if not reader.fieldnames:
            raise ValueError("CSV file needs a header row.")
        text = _records_to_text(records)
    elif suffix in {".html", ".htm"}:
        soup = BeautifulSoup(raw.decode("utf-8", errors="replace"), "html.parser")
        for node in soup(["script", "style", "noscript", "svg"]):
            node.decompose()
        text = soup.get_text("\n", strip=True)
    elif suffix == ".pdf":
        reader = PdfReader(io.BytesIO(raw))
        pages = []
        for page_no, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""
            if page_text.strip():
                pages.append(f"[PDF page {page_no}]\n{page_text}")
        text = "\n\n".join(pages)
    elif suffix == ".docx":
        doc = Document(io.BytesIO(raw))
        blocks = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(blocks)
    else:  # Defensive; allowed extensions above should make this unreachable.
        raise ValueError("Unsupported file type.")

    text = re.sub(r"\x00", "", text).strip()
    if not text:
        raise ValueError("No readable text was extracted. Scanned PDFs need OCR before upload.")
    return text


def chunk_text(text: str, chunk_size: int = 1400, overlap: int = 180) -> list[str]:
    """Chunk mostly on paragraph/line boundaries, with a safe character fallback."""
    clean = re.sub(r"\r\n?", "\n", text).strip()
    if not clean:
        return []
    blocks = [b.strip() for b in re.split(r"\n\s*\n", clean) if b.strip()]
    chunks: list[str] = []
    current = ""
    for block in blocks:
        # Split a single oversized record/paragraph into overlapping windows.
        if len(block) > chunk_size:
            if current:
                chunks.append(current)
                current = ""
            start = 0
            while start < len(block):
                end = min(len(block), start + chunk_size)
                chunks.append(block[start:end])
                if end >= len(block):
                    break
                start = max(start + 1, end - overlap)
            continue
        combined = f"{current}\n\n{block}" if current else block
        if len(combined) <= chunk_size:
            current = combined
        else:
            chunks.append(current)
            tail = current[-overlap:] if overlap else ""
            current = f"{tail}\n\n{block}".strip()
    if current:
        chunks.append(current)
    return chunks
