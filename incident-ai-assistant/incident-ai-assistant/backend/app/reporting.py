from __future__ import annotations
import html
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, ListFlowable, ListItem
from .config import REPORTS_DIR


def _safe_name(title: str) -> str:
    base = re.sub(r"[^a-zA-Z0-9._-]+", "_", title.strip()).strip("._-")
    return (base or "incident_report")[:80]


def _md_lines(markdown: str) -> list[str]:
    return [line.rstrip() for line in markdown.replace("\r\n", "\n").split("\n")]


def _plain_inline(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = re.sub(r"[*`_]", "", text)
    return text.strip()


def create_report_files(report_id: str, title: str, markdown: str) -> dict[str, str]:
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    stem = f"{_safe_name(title)}_{report_id[:8]}"
    md_path = REPORTS_DIR / f"{stem}.md"
    docx_path = REPORTS_DIR / f"{stem}.docx"
    pdf_path = REPORTS_DIR / f"{stem}.pdf"
    md_path.write_text(markdown, encoding="utf-8")
    _write_docx(docx_path, title, markdown)
    _write_pdf(pdf_path, title, markdown)
    return {"md": str(md_path), "docx": str(docx_path), "pdf": str(pdf_path)}


def _write_docx(path: Path, title: str, markdown: str) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(10)
    doc.styles["Normal"].font.color.rgb = RGBColor(38, 48, 61)
    doc.add_heading(title, level=0)
    in_list = False
    for line in _md_lines(markdown):
        stripped = line.strip()
        if not stripped:
            in_list = False
            continue
        if stripped.startswith("# "):
            doc.add_heading(_plain_inline(stripped[2:]), level=1)
        elif stripped.startswith("## "):
            doc.add_heading(_plain_inline(stripped[3:]), level=2)
        elif stripped.startswith("### "):
            doc.add_heading(_plain_inline(stripped[4:]), level=3)
        elif re.match(r"^[-*]\s+", stripped):
            doc.add_paragraph(_plain_inline(re.sub(r"^[-*]\s+", "", stripped)), style="List Bullet")
            in_list = True
        elif re.match(r"^\d+[.)]\s+", stripped):
            doc.add_paragraph(_plain_inline(re.sub(r"^\d+[.)]\s+", "", stripped)), style="List Number")
            in_list = True
        elif stripped.startswith("|"):
            doc.add_paragraph(_plain_inline(stripped.replace("|", "  •  ")), style="Normal")
        elif stripped in {"---", "***"}:
            doc.add_paragraph("—" * 38)
        else:
            doc.add_paragraph(_plain_inline(stripped))
    footer = section.footer.paragraphs[0]
    footer.text = "Generated from indexed internal evidence • Validate findings before external distribution"
    footer.style = doc.styles["Caption"]
    doc.save(path)


def _write_pdf(path: Path, title: str, markdown: str) -> None:
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ReportTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=19, leading=23, alignment=TA_CENTER, textColor=colors.HexColor("#17324d"), spaceAfter=16))
    styles.add(ParagraphStyle(name="ReportH1", parent=styles["Heading1"], fontSize=15, leading=18, spaceBefore=12, spaceAfter=7, keepWithNext=True))
    styles.add(ParagraphStyle(name="ReportH2", parent=styles["Heading2"], fontSize=12, leading=15, spaceBefore=9, spaceAfter=5, keepWithNext=True))
    styles.add(ParagraphStyle(name="ReportBody", parent=styles["BodyText"], fontSize=9.3, leading=13, spaceAfter=6, splitLongWords=True))
    styles.add(ParagraphStyle(name="ReportSmall", parent=styles["BodyText"], fontSize=8, leading=10, textColor=colors.HexColor("#56616f"), spaceAfter=4))
    story = [Paragraph(html.escape(title), styles["ReportTitle"])]
    bullet_buffer: list[str] = []

    def flush_bullets() -> None:
        nonlocal bullet_buffer
        if bullet_buffer:
            items = [ListItem(Paragraph(html.escape(_plain_inline(item)), styles["ReportBody"]), leftIndent=8) for item in bullet_buffer]
            story.append(ListFlowable(items, bulletType="bullet", leftIndent=18, bulletFontSize=7, spaceAfter=5))
            bullet_buffer = []

    for line in _md_lines(markdown):
        stripped = line.strip()
        if not stripped:
            flush_bullets()
            continue
        if re.match(r"^[-*]\s+", stripped):
            bullet_buffer.append(re.sub(r"^[-*]\s+", "", stripped))
            continue
        flush_bullets()
        if stripped.startswith("# "):
            story.append(Paragraph(html.escape(_plain_inline(stripped[2:])), styles["ReportH1"]))
        elif stripped.startswith("## "):
            story.append(Paragraph(html.escape(_plain_inline(stripped[3:])), styles["ReportH1"]))
        elif stripped.startswith("### "):
            story.append(Paragraph(html.escape(_plain_inline(stripped[4:])), styles["ReportH2"]))
        elif stripped in {"---", "***"}:
            story.append(Spacer(1, 5 * mm))
        elif stripped.startswith("|"):
            story.append(Paragraph(html.escape(_plain_inline(stripped.replace("|", "  •  "))), styles["ReportSmall"]))
        else:
            story.append(Paragraph(html.escape(_plain_inline(stripped)), styles["ReportBody"]))
    flush_bullets()

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#667085"))
        canvas.drawString(16 * mm, 10 * mm, "Internal evidence-based draft • Review before distribution")
        canvas.drawRightString(A4[0] - 16 * mm, 10 * mm, f"Page {doc.page}")
        canvas.restoreState()

    pdf = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=16 * mm, leftMargin=16 * mm, topMargin=16 * mm, bottomMargin=18 * mm, title=title, author="AI Incident Assistant")
    pdf.build(story, onFirstPage=footer, onLaterPages=footer)
